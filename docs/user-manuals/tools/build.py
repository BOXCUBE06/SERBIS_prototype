"""Build the SERBIS manuals: Markdown -> .docx (A5) -> PDF -> reading + booklet PDFs.

    python build.py                     both manuals, A4 sheets
    python build.py admin-panel --paper Letter
    python build.py mobile-app --docx-only

Outputs in export/:
    <manual>.docx              editable, same structure as the PDF
    <manual>-reading.pdf       one A5 page per PDF page, for screens
    <manual>-booklet.pdf       imposed for printing (or -booklet-part1.pdf, -part2 ...)

Images: screenshots/annotated/<ID>.png, else screenshots/raw/<ID>.png, else a grey
placeholder, so the manuals build before every screenshot is captured.
"""
import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
import logging

from pypdf import PdfReader, PdfWriter

logging.getLogger("pypdf").setLevel(logging.ERROR)  # LibreOffice PDFs trip a harmless xref warning

import impose_booklet
from common import (ANNOTATED, BOOKLET_MAX_PAGES, EXPORT, LOGOS, MARGIN_BOTTOM, MARGIN_INNER,
                    MARGIN_OUTER, MARGIN_TOP, PAPER, RAW, ROOT, print_width_mm, text_width_mm)

MANUALS = {
    "admin-panel": "SERBIS Admin Panel User Manual",
    "mobile-app": "SERBIS Mobile App User Manual",
}
VERSION = "v1.0, October 2026"
OFFICE = "Municipal Disaster Risk Reduction and Management Office\nMunicipality of Echague, Isabela"
BACK_CONTACT = (
    "Operations Center, Municipal Ground, San Fabian, Echague, Isabela, 3309\n"
    "Hotline: 09176262352 / 09431320604\n"
    "Email: mdrrmo.echague@gmail.com\n"
    "Facebook: Mdrrmo Echague | Page: Echague Rescue"
)
PART_BREAK = "<!-- part-break -->"

FONT = "Arial"
BODY_PT = 10.5
INK = RGBColor(0x1C, 0x2B, 0x24)
HEAD = RGBColor(0x16, 0x48, 0x3A)
RED = RGBColor(0xD6, 0x20, 0x20)


# ---------------------------------------------------------------- Markdown

def parse(md: str) -> list[tuple]:
    """Blocks: (h, level, text) (img, id, path) (legend, [(n, text)]) (ol, n, text)
    (ul, text) (p, text) (part,)."""
    md = md.replace(PART_BREAK, "\n@@PART@@\n")
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    blocks, para = [], []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for line in md.splitlines():
        s = line.strip()
        if not s:
            flush()
        elif s == "@@PART@@":
            flush(); blocks.append(("part",))
        elif m := re.match(r"(#{1,3}) (.+)", s):
            flush(); blocks.append(("h", len(m[1]), m[2]))
        elif m := re.match(r"!\[([^\]]+)\]\(([^)]+)\)", s):
            flush(); blocks.append(("img", m[1], m[2]))
        elif m := re.match(r">\s*\*\*(\d+)\*\*\s*(.+)", s):
            flush()
            if blocks and blocks[-1][0] == "legend":
                blocks[-1][1].append((m[1], m[2]))
            else:
                blocks.append(("legend", [(m[1], m[2])]))
        elif m := re.match(r"(\d+)\. (.+)", s):
            flush(); blocks.append(("ol", m[1], m[2]))
        elif m := re.match(r"- (.+)", s):
            flush(); blocks.append(("ul", m[1]))
        elif para or not blocks or blocks[-1][0] not in ("ol", "ul"):
            para.append(s)
        else:  # continuation of a list item
            kind, *rest = blocks[-1]
            blocks[-1] = (kind, *rest[:-1], rest[-1] + " " + s)
    flush()
    return blocks


def add_inline(par, text: str, size: float = BODY_PT, color=INK) -> None:
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        bold, italic = tok.startswith("**"), tok.startswith("*") and not tok.startswith("**")
        run = par.add_run(tok.strip("*") if (bold or italic) else tok)
        run.bold, run.italic = bold or None, italic or None
        run.font.size, run.font.color.rgb = Pt(size), color


# ---------------------------------------------------------------- docx helpers

def field(par, instr: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, TOC ...)."""
    def fld(kind):
        r = par.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    fld("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    fld("separate")
    par.add_run(placeholder)
    fld("end")


def setup(doc: Document, paper: str) -> None:
    w, h = PAPER[paper]["page"]
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(w), Mm(h)
    # Mirrored: "left" is the inside (fold) edge on every page.
    s.left_margin, s.right_margin = Mm(MARGIN_INNER), Mm(MARGIN_OUTER)
    s.top_margin, s.bottom_margin = Mm(MARGIN_TOP), Mm(MARGIN_BOTTOM)
    s.footer_distance = Mm(7)
    mirror = OxmlElement("w:mirrorMargins")
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    zoom.addnext(mirror) if zoom is not None else settings.insert(0, mirror)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = FONT, Pt(BODY_PT), INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before in (("Heading 1", 17, 0), ("Heading 2", 12.5, 12), ("Heading 3", 11, 8)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = HEAD
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(6)
        st.paragraph_format.keep_with_next = True


def new_section(doc: Document, odd: bool, footer: str | None):
    """footer: "page" for a number, "none" for empty, None to inherit."""
    sec = doc.add_section(WD_SECTION.ODD_PAGE if odd else WD_SECTION.NEW_PAGE)
    if footer:
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0]
        p.text = ""
        if footer == "page":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            field(p, "PAGE", "1")
            for r in p.runs:
                r.font.size = Pt(9.5)
    return sec


def logos(doc: Document, height_mm: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("bagong-pilipinas.png", "echague-seal.png", "echague-rescue.png"):
        path = LOGOS / name
        if path.exists():
            p.add_run().add_picture(str(path), height=Mm(height_mm))
            p.add_run("   ")


def centered(doc: Document, text: str, size: float, bold=False, color=INK, before=0.0):
    for i, line in enumerate(text.split("\n")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before if i == 0 else 0)
        r = p.add_run(line)
        r.font.size, r.bold, r.font.color.rgb = Pt(size), bold, color


def placeholder(shot_id: str, tmp: Path) -> Path:
    """Grey stand-in for a screenshot not captured yet."""
    path = tmp / f"{shot_id}.png"
    size = (390, 844) if shot_id.startswith("MOB") else (1280, 720)
    img = Image.new("RGB", size, (232, 232, 228))
    d = ImageDraw.Draw(img)
    try:
        f = ImageFont.truetype("arialbd.ttf", 48 if size[0] > 400 else 30)
    except OSError:
        f = ImageFont.load_default()
    d.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=(150, 150, 145), width=6)
    d.text((size[0] / 2, size[1] / 2), f"{shot_id}\nnot captured yet", fill=(110, 110, 105),
           font=f, anchor="mm", align="center")
    img.save(path)
    return path


def image_for(shot_id: str, tmp: Path) -> Path:
    for p in (ANNOTATED / f"{shot_id}.png", RAW / f"{shot_id}.png"):
        if p.exists():
            return p
    return placeholder(shot_id, tmp)


# ---------------------------------------------------------------- build

def build_docx(blocks: list, title: str, part: str | None, paper: str, out: Path, tmp: Path) -> None:
    doc = Document()
    setup(doc, paper)

    # Front cover: page 1, no page number.
    logos(doc, 22)
    centered(doc, title, 22, bold=True, color=HEAD, before=60)
    if part:
        centered(doc, part, 14, bold=True, before=6)
    centered(doc, VERSION, 11, before=12)
    centered(doc, OFFICE, 10.5, before=150)

    # Contents on the next right-hand page; numbers from here on.
    new_section(doc, odd=True, footer="page")
    p = doc.add_paragraph()
    r = p.add_run("Contents")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(17), HEAD
    # No \h: hyperlinked entries print in link colours. PDF bookmarks still navigate.
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\z \\u', "Contents are filled in when the PDF is made.")

    def legend(container, items):
        for i, (n, text) in enumerate(items):
            p = container.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Mm(7), Mm(-7)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(7))
            p.paragraph_format.space_after = Pt(1 if i < len(items) - 1 else 8)
            p.paragraph_format.keep_with_next = i < len(items) - 1
            r = p.add_run(n)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(BODY_PT), RED
            add_inline(p, "\t" + text)

    skip = False
    for k, b in enumerate(blocks):
        kind = b[0]
        if skip:  # legend already placed beside its phone screenshot
            skip = False
            continue
        if kind == "h":
            if b[1] == 1:
                new_section(doc, odd=True, footer=None)  # each main section starts on a right-hand page
            doc.add_heading(b[2], level=b[1])
        elif kind == "p":
            add_inline(doc.add_paragraph(), b[1])
        elif kind in ("ol", "ul"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Mm(6), Mm(-6)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(6))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, (f"{b[1]}.\t{b[2]}" if kind == "ol" else f"•\t{b[1]}"))
        elif kind == "img":
            path = image_for(b[1], tmp)
            with Image.open(path) as im:
                width = print_width_mm(im.width, im.height, paper)
                portrait = im.height > im.width
            nxt = blocks[k + 1] if k + 1 < len(blocks) else None
            if portrait and nxt and nxt[0] == "legend":
                # Phone shot: picture left, legend right, so a page holds two.
                table = doc.add_table(rows=1, cols=2)
                table.autofit = False
                pic, cap = table.rows[0].cells
                pic.width, cap.width = Mm(width + 2), Mm(text_width_mm(paper) - width - 2)
                pic.paragraphs[0].add_run().add_picture(str(path), width=Mm(width))
                cap.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cap._element.remove(cap.paragraphs[0]._element)
                legend(cap, nxt[1])
                tr_pr = table.rows[0]._tr.get_or_add_trPr()
                no_split = OxmlElement("w:cantSplit")
                tr_pr.append(no_split)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                skip = True
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(str(path), width=Mm(width))
        elif kind == "legend":
            legend(doc, b[1])

    # Back cover: last page, no page number. Padding goes in front of it later.
    new_section(doc, odd=False, footer="none")
    logos(doc, 16)
    centered(doc, "SERBIS", 16, bold=True, color=HEAD, before=40)
    centered(doc, OFFICE, 10.5, before=6)
    centered(doc, BACK_CONTACT, 10.5, before=12)
    centered(doc, VERSION, 9.5, before=200)
    doc.save(out)


def lo_python() -> Path:
    env = os.environ.get("SOFFICE_PYTHON")
    for p in ([Path(env)] if env else []) + [Path(r"C:\Program Files\LibreOffice\program\python.exe"),
                                             Path(r"C:\Program Files (x86)\LibreOffice\program\python.exe")]:
        if p.exists():
            return p
    sys.exit("LibreOffice not found. Install it, or set SOFFICE_PYTHON to its program\\python.exe.")


def to_pdf(docx: Path, pdf: Path) -> None:
    subprocess.run([str(lo_python()), str(Path(__file__).with_name("lo_export.py")), str(docx), str(pdf)],
                   check=True, timeout=600)


def split_parts(blocks: list) -> list[list]:
    parts, cur = [], []
    for b in blocks:
        if b[0] == "part":
            parts.append(cur)
            cur = []
        else:
            cur.append(b)
    return parts + [cur]


def build(name: str, paper: str, docx_only: bool) -> None:
    title = MANUALS[name]
    blocks = parse((ROOT / f"{name}.md").read_text(encoding="utf-8"))
    parts = split_parts(blocks)
    EXPORT.mkdir(exist_ok=True)
    work = EXPORT / ".build"
    work.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        full_docx = EXPORT / f"{name}.docx"
        build_docx([b for b in blocks if b[0] != "part"], title, None, paper, full_docx, tmp)
        print(f"{name}: {full_docx.name}")
        if docx_only:
            return

        full_pdf = work / f"{name}.pdf"
        to_pdf(full_docx, full_pdf)
        reading = impose_booklet.padded(list(PdfReader(full_pdf).pages))
        out = PdfWriter()
        for pg in reading:
            out.add_page(pg)
        with (EXPORT / f"{name}-reading.pdf").open("wb") as f:
            out.write(f)
        print(f"  {name}-reading.pdf: {len(reading)} pages")

        booklet = EXPORT / f"{name}-booklet.pdf"
        if len(parts) == 1:
            impose_booklet.write_booklets(list(PdfReader(full_pdf).pages), booklet, paper)
            return
        for i, part_blocks in enumerate(parts, 1):
            docx = work / f"{name}-part{i}.docx"
            build_docx(part_blocks, title, f"Part {i} of {len(parts)}", paper, docx, tmp)
            pdf = work / f"{name}-part{i}.pdf"
            to_pdf(docx, pdf)
            pages = list(PdfReader(pdf).pages)
            if len(impose_booklet.padded(pages)) > BOOKLET_MAX_PAGES:
                print(f"  WARNING part {i} is over {BOOKLET_MAX_PAGES} pages: move the part-break.")
            impose_booklet.write_booklets(pages, booklet.with_name(f"{booklet.stem}-part{i}.pdf"), paper)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("manuals", nargs="*", help=f"any of: {', '.join(MANUALS)} (default: all)")
    p.add_argument("--paper", choices=PAPER, default="A4", help="printing sheet; pages are half of it")
    p.add_argument("--docx-only", action="store_true", help="skip PDF export and imposition")
    a = p.parse_args()
    unknown = set(a.manuals) - set(MANUALS)
    if unknown:
        p.error(f"unknown manual: {', '.join(unknown)}")
    for name in a.manuals or MANUALS:
        build(name, a.paper, a.docx_only)
    if not a.docx_only:
        print(impose_booklet.INSTRUCTIONS)
    return 0


if __name__ == "__main__":
    sys.exit(main())
